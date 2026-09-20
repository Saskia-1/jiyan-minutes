#!/usr/bin/env python3
"""Build the synthetic, secret-free formal delivery guide."""

from __future__ import annotations

from io import BytesIO
import os
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def build() -> bytes:
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.top_margin = section.bottom_margin = Inches(0.75)
    section.left_margin = section.right_margin = Inches(0.8)
    document.core_properties.title = "Chat to How-To 交付说明"
    document.core_properties.subject = "macOS 本地聊天导出、搜索、AI 生成与文档交付"
    document.core_properties.author = ""
    document.core_properties.last_modified_by = ""

    for style_name, size in (("Normal", 10.5), ("Title", 25), ("Heading 1", 17), ("Heading 2", 13)):
        style = document.styles[style_name]
        style.font.name = "Hiragino Sans GB"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(26, 32, 44)
    document.styles["Normal"].paragraph_format.space_after = Pt(6)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.add_run("Chat to How-To\n交付说明")
    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.add_run("macOS Apple Silicon · 本地优先 · 2026-09-11").italic = True
    _callout(document, "交付结论", "P1–P3 的本地功能和合成端到端测试已实现；QQ 与微信真实导出、DeepSeek 真实调用仍未实测，不能写成已完成。")

    document.add_heading("1. 交付范围", level=1)
    document.add_paragraph("平台在用户选择并确认后，把受审查的本机来源消息保存到批准目录、校验并导入 SQLite；随后可本地搜索，按独立授权调用 DeepSeek 生成带证据的操作指南，人工编辑不可变版本，并导出 Markdown 或 DOCX。")
    rows = (
        ("P0", "环境与来源", "代码/契约完成；真实 QQ、微信均 blocked / 未实测"),
        ("P1", "导出与搜索", "代码与合成 E2E 完成"),
        ("P2", "DeepSeek 生成", "实现与 mock 完成；真实 API 未实调"),
        ("P3", "文档与交付", "编辑、版本、导出、删除、启动、文档完成"),
    )
    _table(document, ("阶段", "范围", "状态"), rows, (0.65, 1.8, 4.8))

    document.add_heading("2. 快速启动", level=1)
    document.add_paragraph("要求：macOS Apple Silicon、Python 3.11+、Node.js 24.x 与 corepack。")
    _code(document, "./scripts/bootstrap_macos.sh\n./scripts/start_macos.sh\n./scripts/stop_macos.sh")
    document.add_paragraph("初始化会显示工作区外的数据与导出目录并要求输入 YES；已有配置不覆盖。生产前端和 API 由单个进程提供，只监听 http://127.0.0.1:8765。")

    document.add_heading("3. DeepSeek 配置", level=1)
    document.add_paragraph("运行配置默认选择 deepseek-v4-flash；也允许 deepseek-v4-pro。实际 Key 只由用户在本机进程环境显式设置，不写入仓库、SQLite 或前端。")
    _code(document, "export DEEPSEEK_API_KEY=\"<在本机安全设置>\"\n./scripts/start_macos.sh")
    document.add_paragraph("生成前页面展示完整脱敏 user 文本、消息数、字符数、最多调用数与输出上限。只有勾选一次性确认才外发。脱敏不能保证覆盖所有个人信息；用户仍须核对。当前价格来源未配置，费用保持不可用，不凭空估算。取消或超时后 Provider 可能已收到请求并计费。")

    document.add_heading("4. 用户按钮与后台流程", level=1)
    flows = (
        ("首次初始化", "锁文件安装 → 生产构建 → 私有非秘密配置；不连接来源。"),
        ("打开平台", "迁移/恢复 → 同源静态页 → 会话/CSRF → 安全环境状态。"),
        ("保存来源配置", "校验固定回环端点和版本 → SQLite 公开元数据；不读正文。"),
        ("保存/删除 Token", "独立短期确认 → macOS Keychain set/delete；浏览器不回显。"),
        ("读取本机会话列表", "一次 secret_get → 只读适配器列元数据 → 短期不透明句柄。"),
        ("选择聊天/时间", "前端校验 → 后端解析句柄与 UTC 半开范围 → 指纹冻结。"),
        ("导出并建立索引", "一次性授权 → 来源读取 → 原子文件 → 哈希/范围 → 自动导入/快照。"),
        ("搜索本地聊天", "当前快照 FTS 查询；不访问来源、不联网。"),
        ("预览 AI 外发", "范围/排除/预算 → 目标脱敏 → 分块 → 展示完整 user 文本；不调用模型。"),
        ("确认并生成", "一次性授权 → 缓存 → DeepSeek → 严格 JSON → 合并 → 本地证据校验。"),
        ("取消/重试生成", "取消本地 HTTP；可能已有费用。重试同一冻结范围并累计用量记录。"),
        ("查看依据", "只读已授权快照引文；不执行消息中的命令、URL 或提示。"),
        ("创建/编辑/恢复文档", "追加不可变版本；人工版本标为未验证；零模型调用。"),
        ("导出当前版本", "本地 Markdown/python-docx → 原子私有文件 → 回读、大小与 SHA-256。"),
        ("查看删除影响/确认删除", "依赖计数与指纹 → 暂存文件 → DB 事务 → 清理/启动恢复。"),
        ("断开本机来源", "先删除 Keychain Token → 清配置/未用授权；保留快照、文档和文件。"),
    )
    _table(document, ("用户动作", "后台结果"), flows, (1.8, 5.45))
    document.add_paragraph("术语：适配器是本机来源接口转换层；快照是一次导入后冻结的消息集合；幂等表示重复提交不会重复产生任务；Token 是来源 API 秘密；Provider 是外部 AI 服务。每个动作的读取、写入、联网、失败残留、重试和费用详见 docs/how-it-works.md。")

    document.add_heading("5. 隐私与删除边界", level=1)
    for item in (
        "SQLite 默认不加密；数据与导出目录应保持 0700，文件保持 0600，并避免不希望的 iCloud/网盘/备份同步。",
        "断开来源、删除本项目数据、删除受管文件、清理第三方工具缓存、撤销授权、撤销 macOS 权限是六个不同动作。",
        "删除永不向 QQ、微信、QCE、TraceMemo 或 DeepSeek 发送删除请求，也不删除原始聊天。",
        "选择删除文件时先同目录暂存，再提交数据库；崩溃后的 cleanup_pending 会在启动时重试。",
    ):
        document.add_paragraph(item, style="List Bullet")

    document.add_heading("6. 实测与未实测", level=1)
    _table(document, ("证据", "结果"), (
        ("本机", "macOS 26.5.2 / arm64；QQ 7.0.1；微信 4.1.13；Python 3.11.9 arm64；Node 24.19.0 arm64。"),
        ("自动化", "后端/前端单元、严格安全、迁移、合成 API E2E、Playwright 桌面与 360px、构建与类型检查。精确数字见验收报告。"),
        ("QQ", "QCE 未安装；候选路线需复制/修改/重签名运行副本并重新登录，当前边界阻塞。"),
        ("微信", "TraceMemo 未安装；供应链/许可和微信 4.1.13 兼容性未固定；SIP/管理员/内存取钥路径禁止。"),
        ("DeepSeek", "实现和 MockTransport/FakeProvider 已测；没有用真实 Key 调用。"),
    ), (1.35, 5.9))

    document.add_heading("7. 交付文件索引", level=1)
    for item in (
        "README.md：安装、启动、停止、DeepSeek 快速配置与排错。",
        "docs/environment.md：本机安全检测。",
        "docs/source-contracts.md / source-setup.md：来源证据、首次配置和阻碍。",
        "docs/api-setup.md：DeepSeek 外发、错误、用量和费用。",
        "docs/how-it-works.md：逐按钮完整后台流程。",
        "docs/privacy-and-cleanup.md：数据位置和六类清理边界。",
        "docs/acceptance-report.md：精确实测、未实测和 blocker。",
    ):
        document.add_paragraph(item, style="List Bullet")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("Chat to How-To · 本地优先交付说明")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _table(document, headers, rows, widths):
    table = document.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for index, (cell, header) in enumerate(zip(table.rows[0].cells, headers)):
        cell.text = header
        cell.width = Inches(widths[index])
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        cell.paragraphs[0].runs[0].bold = True
        _shade(cell, "DCEAF7")
    for row in rows:
        cells = table.add_row().cells
        for index, (cell, value) in enumerate(zip(cells, row)):
            cell.text = value
            cell.width = Inches(widths[index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    document.add_paragraph()


def _shade(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def _callout(document, heading, body):
    table = document.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    cell = table.cell(0, 0)
    _shade(cell, "EEF6EE")
    paragraph = cell.paragraphs[0]
    paragraph.add_run(heading + "\n").bold = True
    paragraph.add_run(body)
    document.add_paragraph()


def _code(document, value):
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Inches(0.3)
    paragraph.paragraph_format.right_indent = Inches(0.3)
    run = paragraph.add_run(value)
    run.font.name = "Menlo"
    run.font.size = Pt(9)


def main() -> int:
    target = Path(sys.argv[1] if len(sys.argv) > 1 else "docs/Chat-to-How-To-交付说明.docx")
    payload = build()
    temporary = target.with_name("." + target.name + ".tmp")
    descriptor = os.open(temporary, os.O_WRONLY | os.O_CREAT | os.O_TRUNC, 0o600)
    with os.fdopen(descriptor, "wb") as handle:
        handle.write(payload)
        handle.flush()
        os.fsync(handle.fileno())
    os.replace(temporary, target)
    reopened = Document(target)
    joined = "\n".join(paragraph.text for paragraph in reopened.paragraphs)
    if "Chat to How-To" not in joined or "DeepSeek" not in joined:
        raise RuntimeError("delivery DOCX verification failed")
    print(f"created {target.name}: {len(reopened.paragraphs)} paragraphs, {len(reopened.tables)} tables")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
