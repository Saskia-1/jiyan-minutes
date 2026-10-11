"""Private atomic Markdown and DOCX exports from one immutable document version."""

from __future__ import annotations

from datetime import datetime, timezone
from hashlib import sha256
from io import BytesIO
import os
import re
from typing import Literal, Protocol
from uuid import uuid4
from zipfile import ZIP_DEFLATED, ZipFile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_BREAK
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from pydantic import BaseModel, ConfigDict, Field

from chat_to_howto.api.session import digest
from chat_to_howto.domain.documents import DocumentVersion, HowToDocument
from chat_to_howto.domain.generation import TopicSummary, summary_groups
from chat_to_howto.services.artifacts import (
    ApprovedOutputRoot,
    _open_or_create_private_directory,
    _read_small_file,
    _rename_noreplace,
    _write_untracked_private_file,
)
from chat_to_howto.services.evidence_validation import ValidatedEvidence


DocumentFormat = Literal["markdown", "docx"]
SourceMode = Literal["body_only", "with_sources"]


class DocumentArtifact(BaseModel):
    model_config = ConfigDict(frozen=True, extra="forbid", strict=True)

    artifact_id: str
    document_id: str
    version_id: str
    format: DocumentFormat
    source_mode: SourceMode
    filename: str
    byte_size: int = Field(gt=0)
    sha256: str = Field(pattern=r"^[0-9a-f]{64}$")
    created_at: datetime


class _DocumentService(Protocol):
    def get_document(self, document_id: str, *, session_id: str) -> HowToDocument: ...


class _ArtifactRepository(Protocol):
    def record_document_artifact(self, **values: object) -> dict[str, object]: ...


class DocumentExportService:
    def __init__(self, *, document_service: _DocumentService, repository: _ArtifactRepository,
                 approved_roots: dict[str, ApprovedOutputRoot], clock=None) -> None:
        self._documents = document_service
        self._repository = repository
        self._roots = dict(approved_roots)
        self._clock = clock or (lambda: datetime.now(timezone.utc))

    def download(self, document_id: str, *, version_id: str, format: DocumentFormat,
                 source_mode: SourceMode, session_id: str) -> tuple[str, bytes]:
        if format not in {"markdown", "docx"} or source_mode not in {"body_only", "with_sources"}:
            raise ValueError("unsupported document export options")
        document = self._documents.get_document(document_id, session_id=session_id)
        version = next((item for item in document.versions if item.version_id == version_id), None)
        if version is None:
            raise ValueError("document version does not exist")
        payload = (render_markdown(version, document.source_evidence, source_mode=source_mode)
                   if format == "markdown" else render_docx(version, document.source_evidence, source_mode=source_mode))
        return f"{_safe_slug(version.content.title)}.{'md' if format == 'markdown' else 'docx'}", payload

    def export(self, document_id: str, *, version_id: str, format: DocumentFormat,
               source_mode: SourceMode, output_root_id: str, session_id: str) -> DocumentArtifact:
        if format not in {"markdown", "docx"} or source_mode not in {"body_only", "with_sources"}:
            raise ValueError("unsupported document export options")
        root = self._roots.get(output_root_id)
        if root is None:
            raise PermissionError("document output root is not approved")
        document = self._documents.get_document(document_id, session_id=session_id)
        version = next((item for item in document.versions if item.version_id == version_id), None)
        if version is None:
            raise ValueError("document version does not exist")
        payload = (render_markdown(version, document.source_evidence, source_mode=source_mode)
                   if format == "markdown" else render_docx(version, document.source_evidence, source_mode=source_mode))
        now = self._clock().astimezone(timezone.utc)
        artifact_id = f"document-artifact:{uuid4().hex}"
        extension = "md" if format == "markdown" else "docx"
        filename = f"{_safe_slug(version.content.title)}-{artifact_id[-12:]}.{extension}"
        byte_size, digest_value = self._write(root, artifact_id, filename, payload)
        try:
            row = self._repository.record_document_artifact(
                artifact_id=artifact_id,
                session_digest=digest(session_id),
                document_id=document.document_id,
                version_id=version.version_id,
                output_root_id=output_root_id,
                filename=filename,
                format=format,
                source_mode=source_mode,
                byte_size=byte_size,
                sha256_digest=digest_value,
                created_at_utc=now.isoformat(),
            )
        except BaseException:
            self._remove_committed(root, filename)
            raise
        return DocumentArtifact(
            artifact_id=str(row.get("artifact_id", artifact_id)),
            document_id=document.document_id,
            version_id=version.version_id,
            format=format,
            source_mode=source_mode,
            filename=filename,
            byte_size=byte_size,
            sha256=digest_value,
            created_at=now,
        )

    @staticmethod
    def _write(root: ApprovedOutputRoot, artifact_id: str, filename: str, payload: bytes) -> tuple[int, str]:
        root_fd = root.open_bound(); directory_fd = None
        temporary_name = f".{artifact_id.replace(':', '-')}.tmp"
        committed = False
        try:
            directory_fd = _open_or_create_private_directory(root_fd, "documents")
            size, digest_value = _write_untracked_private_file(directory_fd, temporary_name, (payload,))
            if _read_small_file(directory_fd, temporary_name, maximum=16 * 1024 * 1024) != payload:
                raise OSError("document artifact verification failed")
            _rename_noreplace(directory_fd, temporary_name, filename)
            committed = True
            os.fsync(directory_fd)
            if _read_small_file(directory_fd, filename, maximum=16 * 1024 * 1024) != payload:
                raise OSError("document artifact verification failed")
            return size, digest_value
        finally:
            if directory_fd is not None:
                if not committed:
                    try: os.unlink(temporary_name, dir_fd=directory_fd)
                    except FileNotFoundError: pass
                os.close(directory_fd)
            os.close(root_fd)

    @staticmethod
    def _remove_committed(root: ApprovedOutputRoot, filename: str) -> None:
        root_fd = root.open_bound(); directory_fd = None
        try:
            directory_fd = _open_or_create_private_directory(root_fd, "documents")
            os.unlink(filename, dir_fd=directory_fd)
            os.fsync(directory_fd)
        except FileNotFoundError:
            pass
        finally:
            if directory_fd is not None: os.close(directory_fd)
            os.close(root_fd)


def render_markdown(version: DocumentVersion, evidence: tuple[ValidatedEvidence, ...], *, source_mode: SourceMode) -> bytes:
    content = version.content
    lines = [f"# {_md(content.title)}", "", f"版本：{version.ordinal}",
             f"来源：{_origin_label(version.origin)}", f"创建时间：{version.created_at.isoformat()}", ""]
    if content.user_text is not None:
        title, blocks = _edited_text_blocks(content.user_text, fallback_title=content.title)
        lines = [f"# {_md(title)}", ""]
        for kind, text in blocks:
            lines += ([f"## {_md(text)}", ""] if kind == "heading" else [_md(text), ""])
    elif content.summary is not None:
        for label, items in summary_groups(content.summary):
            if items:
                lines += [f"## {_md(label)}", ""]
                for item in items:
                    lines += [(_md(item.text) if isinstance(content.summary, TopicSummary) else f"- {_md(item.text)}"), ""]
                    if source_mode == "with_sources" and isinstance(content.summary, TopicSummary) and content.summary.summary_type != "ideas":
                        lines += [_source_reference(item.source_message_ids, version, evidence), ""]
    else:
        lines += _guide_markdown_body(content)
    if content.user_text is None and content.professional_extensions:
        lines += ["## 专业补充（AI 提供，非原聊天内容）", ""]
        for item in content.professional_extensions:
            lines += [f"### {_md(item.heading)}", "", _md(item.text), ""]
    if source_mode == "with_sources":
        lines += ["## 来源证据", ""]
        for index, item in enumerate(_used_evidence(version, evidence), 1):
            sender = item.sender_name or item.sender_id or "未知发送者"
            lines += [f"### 证据 {index}", "", f"- 发送者：{_md(sender)}",
                      f"- 时间：{item.timestamp.isoformat() if item.timestamp else '时间未知'}", f"- 会话：{_md(item.conversation_id)}",
                      f"- 消息标识：{_md(item.message_id)}", "", _md(item.quote), ""]
    return ("\n".join(lines).rstrip() + "\n").encode("utf-8")


def _guide_markdown_body(content) -> list[str]:
    lines: list[str] = []
    if content.preparations:
        lines += ["## 准备", ""] + [f"- {_md(item.text)}" for item in content.preparations] + [""]
    lines += ["## 步骤", ""]
    for index, step in enumerate(content.steps, 1):
        lines += [f"{index}. {_md(step.instruction)}"]
        if step.location: lines.append(f"   - 位置：{_md(step.location)}")
        if step.parameters: lines.append(f"   - 参数：{'、'.join(_md(item) for item in step.parameters)}")
        if step.conditions: lines.append(f"   - 条件：{'、'.join(_md(item) for item in step.conditions)}")
        if step.warnings: lines.append(f"   - 注意：{'、'.join(_md(item) for item in step.warnings)}")
        if step.expected_result: lines.append(f"   - 预期结果：{_md(step.expected_result)}")
    lines.append("")
    if content.checks:
        lines += ["## 检查", ""] + [f"- {_md(item.text)}" for item in content.checks] + [""]
    if content.missing_information:
        lines += ["## 仍缺信息", ""] + [f"- {_md(item)}" for item in content.missing_information] + [""]
    if content.conflicts:
        lines += ["## 冲突", ""]
        for item in content.conflicts:
            lines += [f"### {_md(item.description)}", "", f"- 说法 A：{_md(item.claim_a)}",
                      f"- 说法 B：{_md(item.claim_b)}",
                      f"- 处理：{_md(item.resolution) if item.resolution else '尚未解析'}", ""]
    return lines


def render_docx(version: DocumentVersion, evidence: tuple[ValidatedEvidence, ...], *, source_mode: SourceMode) -> bytes:
    content = version.content
    document = Document()
    section = document.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = Inches(0.8)
    section.left_margin = section.right_margin = Inches(0.85)
    document.core_properties.author = ""; document.core_properties.last_modified_by = ""
    _set_document_fonts(document)
    document.styles["Normal"].font.size = Pt(11)
    document.styles["Title"].font.size = Pt(24)
    if content.user_text is not None:
        title, blocks = _edited_text_blocks(content.user_text, fallback_title=content.title)
        document.add_paragraph(title, style="Title")
        for kind, text in blocks:
            document.add_paragraph(text, style="Heading 1" if kind == "heading" else "Normal")
    else:
        document.add_paragraph(content.title, style="Title")
        document.add_paragraph(f"版本 {version.ordinal} · {_origin_label(version.origin)} · {version.created_at.isoformat()}")
    if content.user_text is not None:
        pass
    elif content.summary is not None:
        for label, items in summary_groups(content.summary):
            if items:
                document.add_heading(label, level=1)
                for item in items:
                    document.add_paragraph(item.text, style="Normal" if isinstance(content.summary, TopicSummary) else "List Bullet")
                    if source_mode == "with_sources" and isinstance(content.summary, TopicSummary) and content.summary.summary_type != "ideas":
                        document.add_paragraph(_source_reference(item.source_message_ids, version, evidence))
    else:
        _add_guide_docx_body(document, content)
    if content.user_text is None and content.professional_extensions:
        document.add_heading("专业补充（AI 提供，非原聊天内容）", level=1)
        for item in content.professional_extensions:
            document.add_heading(item.heading, level=2)
            document.add_paragraph(item.text)
    if source_mode == "with_sources":
        document.add_heading("来源证据", level=1)
        for index, item in enumerate(_used_evidence(version, evidence), 1):
            document.add_heading(f"证据 {index}", level=2)
            sender = item.sender_name or item.sender_id or "未知发送者"
            document.add_paragraph(f"{sender} · {item.timestamp.isoformat() if item.timestamp else '时间未知'} · {item.conversation_id}")
            document.add_paragraph(item.quote)
            document.add_paragraph(f"消息标识：{item.message_id}")
    buffer = BytesIO(); document.save(buffer)
    return _strip_docx_rsids(buffer.getvalue())


def _set_document_fonts(document: Document) -> None:
    font_name = "Hiragino Sans GB"
    names = ("Normal", "Title", "Heading 1", "Heading 2", "Title Char", "Heading 1 Char", "Heading 2 Char")
    for style_name in names:
        style = document.styles[style_name]
        style.font.name = font_name
        style.font.color.rgb = RGBColor(0, 0, 0)
        fonts = style._element.get_or_add_rPr().get_or_add_rFonts()
        _set_explicit_fonts(fonts, font_name)
        color = style._element.rPr.color
        if color is not None:
            for attribute in ("themeColor", "themeTint", "themeShade"):
                color.attrib.pop(qn(f"w:{attribute}"), None)
    for default in document.styles._element.xpath("./w:docDefaults/w:rPrDefault/w:rPr/w:rFonts"):
        _set_explicit_fonts(default, font_name)


def _set_explicit_fonts(fonts, font_name: str) -> None:
    for attribute in ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme", "csTheme"):
        fonts.attrib.pop(qn(f"w:{attribute}"), None)
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def _add_guide_docx_body(document, content) -> None:
    if content.preparations:
        document.add_heading("准备", level=1)
        for item in content.preparations: document.add_paragraph(item.text, style="List Bullet")
    document.add_heading("步骤", level=1)
    for step in content.steps:
        document.add_paragraph(step.instruction, style="List Number")
        for label, values in (("位置", (step.location,) if step.location else ()), ("参数", step.parameters),
                              ("条件", step.conditions), ("注意", step.warnings),
                              ("预期结果", (step.expected_result,) if step.expected_result else ())):
            if values: document.add_paragraph(f"{label}：{'、'.join(values)}")
    if content.checks:
        document.add_heading("检查", level=1)
        for item in content.checks: document.add_paragraph(item.text, style="List Bullet")
    if content.missing_information:
        document.add_heading("仍缺信息", level=1)
        for item in content.missing_information: document.add_paragraph(item, style="List Bullet")
    if content.conflicts:
        document.add_heading("冲突", level=1)
        for item in content.conflicts:
            document.add_heading(item.description, level=2)
            document.add_paragraph(f"说法 A：{item.claim_a}")
            document.add_paragraph(f"说法 B：{item.claim_b}")
            document.add_paragraph(f"处理：{item.resolution or '尚未解析'}")


def _edited_text_blocks(value: str, *, fallback_title: str) -> tuple[str, tuple[tuple[str, str], ...]]:
    """Parse only explicit edit markers; legacy text never gets guessed headings."""
    lines = value.splitlines()
    first_content = next((index for index, line in enumerate(lines) if line.strip()), None)
    explicit_title = (
        first_content is not None
        and lines[first_content].startswith("# ")
        and not lines[first_content].startswith("## ")
        and bool(lines[first_content][2:].strip())
    )
    title = lines[first_content][2:].strip() if explicit_title and first_content is not None else fallback_title
    start = first_content + 1 if explicit_title and first_content is not None else 0
    blocks: list[tuple[str, str]] = []
    paragraph: list[str] = []

    def flush_paragraph() -> None:
        text = "\n".join(paragraph).strip()
        if text:
            blocks.append(("paragraph", text))
        paragraph.clear()

    for line in lines[start:]:
        if explicit_title and line.startswith("## ") and line[3:].strip():
            flush_paragraph()
            blocks.append(("heading", line[3:].strip()))
        elif not line.strip():
            flush_paragraph()
        else:
            paragraph.append(line)
    flush_paragraph()
    if not explicit_title and blocks and blocks[0] == ("paragraph", fallback_title):
        blocks.pop(0)
    return title, tuple(blocks)


def _used_evidence(version: DocumentVersion, evidence: tuple[ValidatedEvidence, ...]) -> tuple[ValidatedEvidence, ...]:
    identifiers: list[str] = []
    content = version.content
    if content.summary is not None:
        for _label, items in summary_groups(content.summary):
            for item in items: identifiers.extend(item.source_message_ids)
    for item in (*content.preparations, *content.steps, *content.checks): identifiers.extend(item.source_message_ids)
    for item in content.conflicts:
        identifiers.extend(item.claim_a_source_message_ids); identifiers.extend(item.claim_b_source_message_ids)
        identifiers.extend(item.resolution_source_message_ids)
    wanted = set(identifiers)
    return tuple(item for item in evidence if item.message_id in wanted)


def _source_reference(identifiers, version, evidence) -> str:
    numbers = {item.message_id: index for index, item in enumerate(_used_evidence(version, evidence), 1)}
    return f"依据：{len(identifiers)} 条原始消息（" + "、".join(f"证据 {numbers[value]}" for value in identifiers) + "）"


def _origin_label(value: str) -> str:
    return {"ai_generated": "AI 原始版本", "user_edited": "用户编辑版本", "restored": "恢复版本"}[value]


def _md(value: str) -> str:
    return re.sub(r"([\\`*_{}\[\]()#+\-.!|<>])", r"\\\1", value)


def _safe_slug(value: str) -> str:
    cleaned = "".join(character if character.isalnum() else "-" for character in value.strip())
    cleaned = re.sub(r"-+", "-", cleaned).strip("-")[:48]
    return cleaned or "how-to"


def _strip_docx_rsids(payload: bytes) -> bytes:
    source = BytesIO(payload); target = BytesIO()
    with ZipFile(source) as original, ZipFile(target, "w", ZIP_DEFLATED) as cleaned:
        for info in original.infolist():
            value = original.read(info.filename)
            if info.filename.startswith("word/") and info.filename.endswith(".xml"):
                value = re.sub(rb"\s+w:rsid[A-Za-z]*=\"[0-9A-Fa-f]+\"", b"", value)
            cleaned.writestr(info, value)
    return target.getvalue()
