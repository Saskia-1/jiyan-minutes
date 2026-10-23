"""PRD v0.2: raw cleanup preserves editable results, never independent evidence."""
from datetime import timedelta
from io import BytesIO
from datetime import datetime
from threading import Event, Thread

from docx import Document
import pytest

from chat_to_howto.services.document_export import render_markdown, render_docx
from .test_deletion import _service
from .test_import_sessions import Clock, _payload
from chat_to_howto.domain.imports import ImportOptions
from chat_to_howto.repositories.chat import ChatRepository
from chat_to_howto.services.import_sessions import ChatImportService
from chat_to_howto.services.deletion import DeletionService
from chat_to_howto.api.session import digest


def test_raw_cleanup_preserves_versions_and_exports_but_removes_evidence(tmp_path):
    repository, _, documents, document, service, output = _service(tmp_path)
    try:
        exports = output / "exports" / "batch-a"
        exports.parent.mkdir(mode=0o700)
        exports.mkdir(mode=0o700)
        (exports / "messages.jsonl").write_text("synthetic instruction")
        edited = documents.create_version(document.document_id,
            parent_version_id=document.current_version.version_id,
            content=document.current_version.content.model_copy(update={"title": "我的复盘"}),
            change_summary="修改标题", session_id="session-a")
        plan = service.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        assert plan.counts["documents"] == 0
        result = service.confirm(plan.plan_id, fingerprint=plan.fingerprint,
            confirmed=True, session_id="session-a")
        assert result.state == "complete"
        assert not exports.exists()
        assert repository.message_count() == 0
        with pytest.raises(ValueError):
            repository.get_generation_job(document.source_generation_job_id)
        kept = documents.get_document(document.document_id, session_id="session-a")
        assert len(kept.versions) == 2
        assert kept.current_version.content == edited.current_version.content
        assert kept.source_evidence == () and kept.source_deleted
        updated = documents.create_version(document.document_id,
            parent_version_id=kept.current_version.version_id,
            content=kept.current_version.content.model_copy(update={"title": "清理后修改"}),
            change_summary="仍可编辑", session_id="session-a")
        assert updated.current_version.content.title == "清理后修改"
        assert render_markdown(updated.current_version, (), source_mode="body_only")
        assert Document(BytesIO(render_docx(updated.current_version, (), source_mode="body_only")))
    finally:
        repository.close()


def test_raw_cleanup_cannot_leave_managed_raw_files(tmp_path):
    repository, _, _, _, service, _ = _service(tmp_path)
    try:
        with pytest.raises(ValueError):
            service.create_plan(target_kind="batch", target_id="batch-a",
                include_managed_files=False, preserve_documents=True, session_id="session-a")
    finally:
        repository.close()


def test_with_sources_exports_disclose_persistent_quotes(tmp_path):
    repository, _, _, document, _, _ = _service(tmp_path)
    notice = "本文件包含原始对话摘录；删除原文不会删除本文件中的摘录。"
    try:
        assert notice in render_markdown(document.current_version, document.source_evidence,
            source_mode="with_sources").decode()
        word = Document(BytesIO(render_docx(document.current_version, document.source_evidence,
            source_mode="with_sources")))
        assert notice in "\n".join(p.text for p in word.paragraphs)
    finally:
        repository.close()


def test_expiry_is_seven_days_from_confirmation_and_cleanup_is_idempotent(tmp_path):
    from chat_to_howto.services.retention import RetentionService
    repository = ChatRepository(tmp_path / "chat.sqlite3")
    clock = Clock()
    imports = ChatImportService(repository=repository, clock=clock)
    try:
        preview = imports.create_preview(filename="chat.csv", payload=_payload(),
            origin="imported_file", platform_hint="qq",
            options=ImportOptions(timezone="Asia/Shanghai"), session_id="session-a")
        clock.now += timedelta(minutes=1)
        committed = imports.commit_preview(preview.preview_id,
            preview_fingerprint=preview.preview_fingerprint, confirmed=True, session_id="session-a")
        assert committed.raw_expires_at_utc == clock.now + timedelta(days=7)
        deletion = DeletionService(repository=repository, approved_roots={}, clock=clock)
        retention = RetentionService(repository=repository, deletion=deletion, clock=clock)
        clock.now += timedelta(days=7) - timedelta(seconds=1)
        retention.sweep()
        assert repository.message_count() == 1
        clock.now += timedelta(seconds=1)
        retention.sweep()
        assert repository.message_count() == 0
        assert imports.recover_committed(session_id="session-a").imports == ()
        assert retention.sweep() == 0
    finally:
        imports.close()
        repository.close()


def test_partial_unsearchable_raw_still_expires(tmp_path):
    from chat_to_howto.services.retention import RetentionService
    repository, _, _, _, deletion, _ = _service(tmp_path)
    try:
        created = datetime.fromisoformat(repository.raw_retention_targets()[0]["created_at_utc"])
        with repository._transaction():
            repository._connection.execute("UPDATE chat_export_jobs SET searchable=0,state='partial' WHERE job_id='export-job'")
        retention = RetentionService(repository=repository, deletion=deletion,
            clock=lambda: created + timedelta(days=8))
        retention.sweep()
        assert repository.message_count() == 0
    finally:
        repository.close()


def test_raw_cleanup_removes_crash_export_temp(tmp_path):
    repository, _, _, _, service, output = _service(tmp_path)
    try:
        exports = output / "exports"; exports.mkdir(mode=0o700)
        temporary = exports / (".batch-a.tmp-" + "a" * 32)
        temporary.mkdir(mode=0o700)
        (temporary / "messages.jsonl").write_text("synthetic instruction")
        plan = service.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        service.confirm(plan.plan_id, fingerprint=plan.fingerprint, confirmed=True, session_id="session-a")
        assert not temporary.exists()
    finally:
        repository.close()


def test_deleted_raw_cannot_dispatch_from_an_in_memory_scope(tmp_path, monkeypatch):
    from chat_to_howto.providers.fake import FakeProvider
    from .test_generation_jobs import _repository_with_snapshot, _orchestrator, _selection, NOW
    repository, _ = _repository_with_snapshot(tmp_path)
    provider = FakeProvider()
    generator = _orchestrator(repository, provider)
    entered, release = Event(), Event()
    worker = None
    outcomes = []
    try:
        preview = generator.create_preview(_selection(), session_id="session-a")
        consent = generator.issue_consent(preview.preview_id, session_id="session-a", confirmed=True)
        job = generator.start_job(consent.consent_id, session_id="session-a", idempotency_key="cleanup-race")
        original = repository.create_generation_request
        def pause(**kwargs):
            value = original(**kwargs)
            entered.set(); assert release.wait(5)
            return value
        monkeypatch.setattr(repository, "create_generation_request", pause)
        def run():
            try: outcomes.append(generator.run_job(job.job_id))
            except Exception as error: outcomes.append(type(error).__name__)
        worker = Thread(target=run); worker.start(); assert entered.wait(5)
        from chat_to_howto.services.artifacts import ApprovedOutputRoot
        output = tmp_path / "approved"; output.mkdir(mode=0o700)
        source_root = tmp_path / "repo"; source_root.mkdir()
        root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=source_root)
        deletion = DeletionService(repository=repository, approved_roots={"approved-root": root}, clock=lambda: NOW)
        plan = deletion.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        deletion.confirm(plan.plan_id, fingerprint=plan.fingerprint, confirmed=True, session_id="session-a")
        release.set(); worker.join(5)
        assert not worker.is_alive()
        assert provider.call_count == 0
        assert len(outcomes) == 1 and not isinstance(outcomes[0], str)
    finally:
        release.set()
        if worker is not None: worker.join(5)
        repository.close()


def test_ocr_crash_images_are_removed_at_startup(tmp_path):
    from chat_to_howto.repositories.content import ContentRepository
    from chat_to_howto.services.content_imports import ContentImportService
    from .test_documents import NOW
    repository = ChatRepository(tmp_path / "chat.sqlite3")
    images = tmp_path / "images"; images.mkdir(mode=0o700)
    temporary = images / ("temp-" + "a" * 32 + ".png")
    temporary.write_bytes(b"synthetic raw screenshot")
    orphan = images / ("b" * 32 + ".jpg"); orphan.write_bytes(b"uncommitted screenshot")
    unrelated = images / "unrecognized-file"; unrelated.write_bytes(b"retain")
    service = ContentImportService(repository=ContentRepository(repository), image_root=images,
        clock=lambda: NOW, ocr=lambda _: ())
    try:
        service.reconcile_orphan_images()
        assert not temporary.exists() and not orphan.exists()
        assert unrelated.read_bytes() == b"retain"
    finally:
        service.close(); repository.close()


def test_legacy_unowned_export_metadata_does_not_break_startup(tmp_path):
    from chat_to_howto.services.retention import RetentionService
    from test_migrations import _create_legacy_v1_database
    from .test_documents import NOW
    path = tmp_path / "legacy.sqlite3"
    _create_legacy_v1_database(path)
    repository = ChatRepository(path)
    try:
        assert repository.retired_export_temp_owners() == ()
        deletion = DeletionService(repository=repository, approved_roots={}, clock=lambda: NOW)
        assert RetentionService(repository=repository, deletion=deletion, clock=lambda: NOW).sweep() == 0
    finally:
        repository.close()


def test_unbound_but_committed_raw_is_cleaned_after_deadline(tmp_path):
    from chat_to_howto.services.retention import RetentionService
    repository, _, documents, document, deletion, output = _service(tmp_path)
    try:
        created = datetime.fromisoformat(repository.raw_retention_targets()[0]["created_at_utc"])
        exports = output / "exports"; exports.mkdir(mode=0o700)
        batch = exports / "batch-a"; batch.mkdir(mode=0o700)
        (batch / "messages.jsonl").write_text("synthetic instruction")
        with repository._transaction():
            repository._connection.execute("UPDATE chat_export_jobs SET batch_id=NULL,snapshot_id=NULL,searchable=0,state='failed' WHERE job_id='export-job'")
        RetentionService(repository=repository, deletion=deletion,
            clock=lambda: created + timedelta(days=8)).sweep()
        assert repository.message_count() == 0
        assert not batch.exists()
        assert documents.get_document(document.document_id, session_id="session-a").source_deleted
    finally:
        repository.close()


def test_retry_timestamp_does_not_extend_raw_deadline(tmp_path):
    from chat_to_howto.services.retention import RetentionService
    repository, _, _, _, deletion, _ = _service(tmp_path)
    try:
        created = datetime.fromisoformat(repository.raw_retention_targets()[0]["created_at_utc"])
        with repository._transaction():
            repository._connection.execute("UPDATE chat_export_jobs SET updated_at_utc=? WHERE job_id='export-job'", ((created + timedelta(days=6)).isoformat(),))
        RetentionService(repository=repository, deletion=deletion,
            clock=lambda: created + timedelta(days=7)).sweep()
        assert repository.message_count() == 0
    finally:
        repository.close()


def test_cleanup_cannot_resurrect_mixed_cache(tmp_path, monkeypatch):
    from datetime import timezone
    from chat_to_howto.api.session import digest
    from chat_to_howto.providers.fake import FakeProvider
    from chat_to_howto.repositories.content import ContentRepository
    from chat_to_howto.services.generation_cache import GenerationCacheSettings
    from chat_to_howto.services.generation_jobs import GenerationJobOrchestrator
    from chat_to_howto.services.mixed_search import MixedSearchService
    from .test_mixed_generation_scope import _selection
    now = datetime(2026, 10, 14, tzinfo=timezone.utc)
    repository = ChatRepository(tmp_path / "chat.sqlite3")
    content = ContentRepository(repository)
    unit_id = "content:v1:note:" + "d" * 64
    content.save_batch(batch_id="note-batch", session_digest=digest("owner"), created_at_utc=now.isoformat(),
        images=[], units=[{"unit_id": unit_id, "image_id": None, "kind": "note_passage", "group_id": "note-1",
            "source_order": 0, "text": "synthetic unique raw note", "sender_name": None, "sent_at": None}])
    generator = GenerationJobOrchestrator(repository=repository, provider_factory=lambda _: FakeProvider(),
        cache_settings=GenerationCacheSettings(prompt_version="synthetic", output_schema_version="synthetic",
            chunk_strategy_version="synthetic", max_chunk_input_tokens=8000, provider_behavior_revision="synthetic"),
        mixed_search=MixedSearchService(repository, content), clock=lambda: now)
    entered, release, in_put = Event(), Event(), Event()
    cleanup_done = Event()
    errors, workers = [], []
    try:
        preview = generator.create_mixed_preview(_selection("note-batch", (unit_id,)), session_id="owner")
        consent = generator.issue_consent(preview.preview_id, session_id="owner", confirmed=True)
        job = generator.start_job(consent.consent_id, session_id="owner", idempotency_key="cache-race")
        original_put, original_units = repository.put_generation_cache, ContentRepository.list_units
        def put(**kwargs):
            in_put.set(); return original_put(**kwargs)
        def units(self, *args, **kwargs):
            result = original_units(self, *args, **kwargs)
            if in_put.is_set(): entered.set(); assert release.wait(5)
            return result
        monkeypatch.setattr(repository, "put_generation_cache", put)
        monkeypatch.setattr(ContentRepository, "list_units", units)
        def generate():
            try: generator.run_job(job.job_id)
            except Exception as error: errors.append(type(error).__name__)
        deletion = DeletionService(repository=repository, approved_roots={}, clock=lambda: now,
            cancel_generation_jobs=generator.revoke_jobs)
        def clean():
            try:
                plan = deletion.create_plan(target_kind="content_batch", target_id="note-batch",
                    include_managed_files=True, preserve_documents=True, session_id="owner")
                assert deletion.confirm(plan.plan_id, fingerprint=plan.fingerprint,
                    confirmed=True, session_id="owner").state == "complete"
            except Exception as error: errors.append(type(error).__name__)
            finally: cleanup_done.set()
        worker = Thread(target=generate); workers.append(worker); worker.start(); assert entered.wait(5)
        cleaner = Thread(target=clean); workers.append(cleaner); cleaner.start()
        cleanup_done.wait(0.2); release.set()
        for worker in workers: worker.join(5); assert not worker.is_alive()
        assert not errors
        assert repository._connection.execute("SELECT COUNT(*) FROM content_units").fetchone()[0] == 0
        assert repository._connection.execute("SELECT COUNT(*) FROM chat_generation_cache_v2").fetchone()[0] == 0
    finally:
        release.set()
        for worker in workers: worker.join(5)
        repository.close()


def test_expired_raw_is_revoked_if_managed_cleanup_fails(tmp_path, monkeypatch):
    from chat_to_howto.services.retention import RetentionService
    from chat_to_howto.api.session import digest
    repository, _, documents, document, deletion, _ = _service(tmp_path)
    try:
        created = datetime.fromisoformat(repository.raw_retention_targets()[0]["created_at_utc"])
        def fail(*args, **kwargs):
            raise PermissionError("synthetic unavailable output")
        monkeypatch.setattr(deletion, "_stage_files", fail)
        retention = RetentionService(repository=repository, deletion=deletion,
            clock=lambda: created + timedelta(days=7))
        assert retention.sweep() == 0 and retention.pending_count > 0
        assert repository.message_count() == 1  # Disk failures are retryable.
        with pytest.raises(PermissionError):
            repository.generation_snapshot_messages(session_digest=digest("session-a"), snapshot_id="snapshot-a")
        assert documents.get_document(document.document_id, session_id="session-a").source_evidence == ()
    finally:
        repository.close()


def test_v14_upgrade_preserves_nonempty_documents_versions_and_artifacts(tmp_path):
    import chat_to_howto.migrations as migrations
    from chat_to_howto.api.session import digest
    from chat_to_howto.services.documents import DocumentService
    repository, _, _, document, _, _ = _service(tmp_path)
    expected = document.current_version.content
    repository.record_document_artifact(artifact_id="artifact:retained", session_digest=digest("session-a"),
            document_id=document.document_id, version_id=document.current_version.version_id,
            output_root_id="approved-root", filename="retained.md", format="markdown", source_mode="body_only",
            byte_size=4, sha256_digest="a" * 64, created_at_utc=document.created_at.isoformat())
    # Recreate the exact v14 document table, including its generation-job FK,
    # then run the real migration over nonempty dependent tables.
    connection = repository._connection
    connection.execute("PRAGMA foreign_keys=OFF")
    connection.execute(migrations._MIGRATION_7[0].replace("CREATE TABLE chat_documents (", "CREATE TABLE chat_documents_v14 ("))
    connection.execute("INSERT INTO chat_documents_v14 SELECT document_id,session_digest,source_generation_job_id,current_version_id,source_result_json,created_at_utc,updated_at_utc FROM chat_documents")
    connection.execute("DROP TABLE chat_documents")
    connection.execute("ALTER TABLE chat_documents_v14 RENAME TO chat_documents")
    connection.execute("CREATE INDEX idx_chat_documents_session ON chat_documents(session_digest,updated_at_utc)")
    connection.execute("DROP TABLE chat_raw_retirements")
    connection.execute("ALTER TABLE chat_deletion_plans DROP COLUMN preserve_documents")
    connection.execute("DELETE FROM chat_schema_migrations WHERE version=15")
    repository.close()
    repository = ChatRepository(tmp_path / "chat.sqlite3")
    try:
        kept = DocumentService(repository=repository).get_document(document.document_id, session_id="session-a")
        assert kept.current_version.content == expected and not kept.source_deleted
        assert len(kept.versions) == 1
        assert len(repository.document_artifacts(document.document_id, session_digest=digest("session-a"))) == 1
        assert repository._connection.execute("PRAGMA foreign_key_check").fetchall() == []
    finally:
        repository.close()


def test_startup_restores_files_staged_before_database_commit(tmp_path):
    repository, _, _, _, service, output = _service(tmp_path)
    try:
        exports = output / "exports"; exports.mkdir(mode=0o700)
        batch = exports / "batch-a"; batch.mkdir(mode=0o700)
        (batch / "messages.jsonl").write_text("synthetic raw")
        plan = service.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        impact = repository.deletion_impact(target_kind="batch", target_id="batch-a",
            session_digest=digest("session-a"), preserve_documents=True)
        service._stage_files(plan.plan_id, impact["managed_files"])
        assert not batch.exists() and repository.interrupted_deletion_staging()
        service.reconcile_cleanup()
        assert (batch / "messages.jsonl").read_text() == "synthetic raw"
        assert not repository.interrupted_deletion_staging()
        assert repository.message_count() == 1
    finally:
        repository.close()


def test_uncommitted_published_export_is_removed_after_deadline(tmp_path):
    repository, _, _, _, deletion, output = _service(tmp_path)
    try:
        created = datetime.fromisoformat(repository.raw_retention_targets()[0]["created_at_utc"])
        exports = output / "exports"; exports.mkdir(mode=0o700)
        orphan = exports / "published-orphan"; orphan.mkdir(mode=0o700)
        (orphan / "messages.jsonl").write_text("synthetic raw")
        with repository._transaction():
            repository._connection.execute("UPDATE chat_export_jobs SET planned_batch_id='published-orphan',state='failed' WHERE job_id='export-job'")
        deletion._clock = lambda: created + timedelta(days=7) - timedelta(seconds=1)
        assert deletion.reconcile_export_temporaries() == 0 and orphan.exists()
        deletion._clock = lambda: created + timedelta(days=7)
        assert deletion.reconcile_export_temporaries() == 0 and not orphan.exists()
    finally:
        repository.close()


def test_cleanup_cancels_inflight_provider_without_recreating_process_records(tmp_path):
    from chat_to_howto.providers.fake import FakeProvider
    from .test_generation_jobs import _repository_with_snapshot, _orchestrator, _selection, NOW
    repository, _ = _repository_with_snapshot(tmp_path)
    entered = Event()
    class WaitingProvider(FakeProvider):
        async def generate(self, request, cancellation):
            entered.set()
            return await super().generate(request, cancellation)
    provider = WaitingProvider(mode="cancellation")
    generator = _orchestrator(repository, provider)
    outcomes, errors = [], []
    worker = None
    try:
        preview = generator.create_preview(_selection(), session_id="session-a")
        consent = generator.issue_consent(preview.preview_id, session_id="session-a", confirmed=True)
        job = generator.start_job(consent.consent_id, session_id="session-a", idempotency_key="inflight-cleanup")
        def run():
            try: outcomes.append(generator.run_job(job.job_id))
            except Exception as error: errors.append(type(error).__name__)
        worker = Thread(target=run); worker.start(); assert entered.wait(5)
        from chat_to_howto.services.artifacts import ApprovedOutputRoot
        output = tmp_path / "approved"; output.mkdir(mode=0o700)
        source_root = tmp_path / "repo"; source_root.mkdir()
        root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=source_root)
        deletion = DeletionService(repository=repository, approved_roots={"approved-root": root}, clock=lambda: NOW,
            cancel_generation_jobs=generator.revoke_jobs)
        plan = deletion.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        deletion.confirm(plan.plan_id, fingerprint=plan.fingerprint, confirmed=True, session_id="session-a")
        worker.join(5)
        assert not worker.is_alive() and not errors
        assert len(outcomes) == 1 and outcomes[0].state == "cancelled" and outcomes[0].result is None
        assert repository._connection.execute("SELECT COUNT(*) FROM chat_generation_requests").fetchone()[0] == 0
    finally:
        if worker is not None and worker.is_alive():
            generator.revoke_jobs((job.job_id,)); worker.join(5)
        repository.close()


def test_multichunk_job_crossing_deadline_stops_and_cleans_without_http(tmp_path):
    from chat_to_howto.providers.fake import FakeProvider
    from chat_to_howto.services.generation_jobs import GenerationJobOrchestrator
    from chat_to_howto.services.retention import RetentionService
    from chat_to_howto.services.artifacts import ApprovedOutputRoot
    from .test_generation_jobs import _repository_with_snapshot, _cache_settings, _selection, NOW
    repository, _ = _repository_with_snapshot(tmp_path)
    clock = Clock(); clock.now = NOW
    class AdvancingProvider(FakeProvider):
        async def generate(self, request, cancellation):
            response = await super().generate(request, cancellation)
            clock.now = NOW + timedelta(days=7)
            return response
    provider = AdvancingProvider()
    generator = GenerationJobOrchestrator(repository=repository, provider_factory=lambda _: provider,
        cache_settings=_cache_settings().model_copy(update={"max_chunk_input_tokens": 180}), clock=clock)
    output = tmp_path / "approved"; output.mkdir(mode=0o700)
    source_root = tmp_path / "repo"; source_root.mkdir()
    root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=source_root)
    deletion = DeletionService(repository=repository, approved_roots={"approved-root": root}, clock=clock,
        cancel_generation_jobs=generator.revoke_jobs)
    retention = RetentionService(repository=repository, deletion=deletion, clock=clock)
    generator._retention_sweep = retention.sweep
    try:
        selection = _selection().model_copy(update={"summary_type": "ideas",
            "end_time": NOW - timedelta(days=1)})
        preview = generator.create_preview(selection, session_id="session-a")
        assert preview.provider_calls > 1
        consent = generator.issue_consent(preview.preview_id, session_id="session-a", confirmed=True)
        job = generator.start_job(consent.consent_id, session_id="session-a", idempotency_key="cross-deadline")
        result = generator.run_job(job.job_id)
        assert provider.call_count == 1
        assert result.state == "cancelled" and result.result is None
        assert repository.message_count() == 0
        assert repository._connection.execute("SELECT COUNT(*) FROM chat_generation_cache").fetchone()[0] == 0
    finally:
        repository.close()


def test_raw_cleanup_cannot_resurrect_preview_from_a_captured_scope(tmp_path):
    from chat_to_howto.providers.fake import FakeProvider
    from chat_to_howto.services.artifacts import ApprovedOutputRoot
    from .test_generation_jobs import _repository_with_snapshot, _orchestrator, _selection, NOW
    repository, _ = _repository_with_snapshot(tmp_path)
    generator = _orchestrator(repository, FakeProvider())
    output = tmp_path / "approved"; output.mkdir(mode=0o700)
    source_root = tmp_path / "repo"; source_root.mkdir()
    root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=source_root)
    deletion = DeletionService(repository=repository, approved_roots={"approved-root": root}, clock=lambda: NOW)
    try:
        scope = generator._scope_builder.build(_selection(), session_id="session-a")
        plan = deletion.create_plan(target_kind="batch", target_id="batch-a",
            include_managed_files=True, preserve_documents=True, session_id="session-a")
        deletion.confirm(plan.plan_id, fingerprint=plan.fingerprint, confirmed=True, session_id="session-a")
        with pytest.raises(PermissionError):
            generator._create_preview_from_scope(scope, session_id="session-a")
        assert repository._connection.execute("SELECT COUNT(*) FROM chat_generation_previews").fetchone()[0] == 0
    finally:
        repository.close()


def test_partial_batch_without_sealed_snapshot_still_removes_base_messages(tmp_path):
    from .test_generation_jobs import _repository_with_snapshot, NOW
    from chat_to_howto.services.artifacts import ApprovedOutputRoot
    from chat_to_howto.services.retention import RetentionService
    repository, _ = _repository_with_snapshot(tmp_path)
    output = tmp_path / "approved"; output.mkdir(mode=0o700)
    source_root = tmp_path / "repo"; source_root.mkdir()
    root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=source_root)
    deletion = DeletionService(repository=repository, approved_roots={"approved-root": root}, clock=lambda: NOW + timedelta(days=7))
    try:
        # Reproduce a committed import that failed before sealing a snapshot.
        with repository._transaction():
            repository._connection.execute("UPDATE chat_export_jobs SET snapshot_id=NULL,searchable=0,state='failed'")
            repository._connection.execute("DROP TRIGGER chat_snapshot_messages_immutable_delete")
            repository._connection.execute("DROP TRIGGER chat_snapshots_immutable_delete")
            repository._connection.execute("DELETE FROM chat_snapshot_messages")
            repository._connection.execute("DELETE FROM chat_snapshots")
        RetentionService(repository=repository, deletion=deletion, clock=lambda: NOW + timedelta(days=7)).sweep()
        assert repository.message_count() == 0
    finally:
        repository.close()
