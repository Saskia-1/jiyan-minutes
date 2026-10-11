"""Markdown and DOCX exports share one local, plain-text view model."""

from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import os
from pathlib import Path
import stat
from zipfile import ZipFile

from docx import Document
from docx.oxml.ns import qn

from chat_to_howto.domain.documents import DocumentVersion, HowToDocument
from chat_to_howto.domain.generation import Check, Conflict, GenerationResult, Preparation, Step, ProfessionalExtension, TopicSection, TopicSummary
from chat_to_howto.services.artifacts import ApprovedOutputRoot
from chat_to_howto.services.document_export import (
    DocumentExportService,
    render_docx,
    render_markdown,
)
from chat_to_howto.services.evidence_validation import ValidatedEvidence


NOW = datetime(2026, 9, 6, tzinfo=timezone.utc)
MESSAGE_ID = "message:v1:" + "c" * 64


def _document() -> HowToDocument:
    content = GenerationResult(
        status="complete",
        title="中文 / English .. 特殊 # 标题",
        preparations=(Preparation(text="准备 <设备> & 账户", source_message_ids=(MESSAGE_ID,)),),
        steps=(Step(instruction="打开 **设置** 并选择 A/B", location="设置 > 本地", parameters=("模式=安全",),
            conditions=("如果已登录",), warnings=("不要泄露密钥",), expected_result="显示 ✓ 完成",
            source_message_ids=(MESSAGE_ID,)),),
        checks=(Check(text="确认结果包含 中文 & English", source_message_ids=(MESSAGE_ID,)),),
        missing_information=("仍需确认端口",),
        conflicts=(Conflict(description="两个说法冲突", claim_a="使用 A", claim_a_source_message_ids=(MESSAGE_ID,),
            claim_b="使用 B", claim_b_source_message_ids=(MESSAGE_ID,), resolution=None, resolution_source_message_ids=()),),
    )
    version = DocumentVersion(
        document_id="document:p3", version_id="document-version:p3", ordinal=1,
        origin="ai_generated", evidence_status="verified", parent_version_id=None,
        restored_from_version_id=None, source_generation_job_id="generation-job:p2",
        change_summary="", created_at=NOW, content=content,
    )
    evidence = ValidatedEvidence(message_id=MESSAGE_ID, conversation_id="conversation-a", timestamp=NOW,
        sender_id="sender-a", sender_name="合成发送者", quote="完整证据 <tag> & **literal**")
    return HowToDocument(document_id="document:p3", source_generation_job_id="generation-job:p2",
        source_snapshot_id="snapshot:p1", created_at=NOW, updated_at=NOW,
        current_version=version, versions=(version,), source_evidence=(evidence,))


def _all_docx_text(payload: bytes) -> str:
    document = Document(BytesIO(payload))
    return "\n".join(paragraph.text for paragraph in document.paragraphs)


def _edited_version(user_text: str, *, title: str = "已保存标题") -> DocumentVersion:
    original = _document().current_version
    return original.model_copy(update={
        "origin": "user_edited",
        "evidence_status": "unverified",
        "content": original.content.model_copy(update={"title": title, "user_text": user_text}),
    })


def test_professional_extension_is_labeled_and_not_added_to_original_evidence():
    doc = _document()
    content = doc.current_version.content.model_copy(update={"professional_extensions": (
        ProfessionalExtension(heading="机制说明", text="这是 AI 的解释，不是参与者原话。"),)})
    version = doc.current_version.model_copy(update={"content": content})
    for text in (render_markdown(version, doc.source_evidence, source_mode="with_sources").decode(),
                 _all_docx_text(render_docx(version, doc.source_evidence, source_mode="with_sources"))):
        assert "专业补充" in text and "这是 AI 的解释" in text
        assert "完整证据" in text
    assert "AI 的解释" not in doc.source_evidence[0].quote


def test_markdown_is_utf8_structured_and_escapes_user_markdown_as_literal_text() -> None:
    document = _document()
    payload = render_markdown(document.current_version, document.source_evidence, source_mode="with_sources")
    text = payload.decode("utf-8")

    assert text.startswith("# 中文 / English \\.\\. 特殊 \\# 标题")
    assert "## 步骤" in text
    assert "1. 打开 \\*\\*设置\\*\\* 并选择 A/B" in text
    assert "## 来源证据" in text
    assert "完整证据 \\<tag\\> & \\*\\*literal\\*\\*" in text
    assert "DEEPSEEK_API_KEY" not in text
    assert "Return JSON only" not in text


def test_docx_has_readable_hierarchy_multilingual_text_and_no_private_metadata() -> None:
    document = _document()
    payload = render_docx(document.current_version, document.source_evidence, source_mode="with_sources")
    text = _all_docx_text(payload)

    for expected in ("中文 / English .. 特殊 # 标题", "准备 <设备> & 账户", "打开 **设置** 并选择 A/B", "来源证据", "完整证据 <tag> & **literal**"):
        assert expected in text
    with ZipFile(BytesIO(payload)) as archive:
        core = archive.read("docProps/core.xml")
        assert b"Saskia" not in core
        assert b"/Users/" not in b"".join(archive.read(name) for name in archive.namelist() if name.endswith(".xml"))
        assert b"rsid" not in archive.read("word/document.xml")


def test_marked_user_edit_renders_real_markdown_and_docx_hierarchy_without_visible_markers() -> None:
    version = _edited_version(
        "# 编辑后的标题\n\n## 1. 核心主题：先表达概括\n\n第一段。\n\n第二段。",
        title="编辑后的标题",
    )

    markdown = render_markdown(version, (), source_mode="body_only").decode()
    assert markdown == "# 编辑后的标题\n\n## 1\\. 核心主题：先表达概括\n\n第一段。\n\n第二段。\n"

    document = Document(BytesIO(render_docx(version, (), source_mode="body_only")))
    assert [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs] == [
        ("Title", "编辑后的标题"),
        ("Heading 1", "1. 核心主题：先表达概括"),
        ("Normal", "第一段。"),
        ("Normal", "第二段。"),
    ]
    assert "#" not in _all_docx_text(render_docx(version, (), source_mode="body_only"))


def test_ideas_export_keeps_internal_sources_but_omits_per_idea_citation_lines() -> None:
    original = _document()
    content = original.current_version.content.model_copy(update={
        "steps": (), "preparations": (), "checks": (), "conflicts": (), "missing_information": (),
        "summary": TopicSummary(summary_type="ideas", sections=(
            TopicSection(heading="1. 场景提醒", text="在特定位置提醒用户。", source_message_ids=(MESSAGE_ID,)),
            TopicSection(heading="2. English 2 与中文标点：另一构思", text="根据既有记录生成摘要。", source_message_ids=(MESSAGE_ID,)),
        )),
    })
    version = original.current_version.model_copy(update={"content": content})

    markdown = render_markdown(version, original.source_evidence, source_mode="with_sources").decode()
    docx = Document(BytesIO(render_docx(version, original.source_evidence, source_mode="with_sources")))

    assert "依据：" not in markdown
    assert "依据：" not in "\n".join(paragraph.text for paragraph in docx.paragraphs)
    assert "## 来源证据" in markdown
    assert "完整证据" in markdown
    assert any(paragraph.text == "完整证据 <tag> & **literal**" for paragraph in docx.paragraphs)
    assert [paragraph.style.name for paragraph in docx.paragraphs if paragraph.text.startswith(("1. 场景", "2. English"))] == [
        "Heading 1", "Heading 1",
    ]


def test_docx_title_and_mixed_language_idea_headings_have_one_explicit_font_contract() -> None:
    original = _document()
    content = original.current_version.content.model_copy(update={
        "steps": (), "preparations": (), "checks": (), "conflicts": (), "missing_information": (),
        "summary": TopicSummary(summary_type="ideas", sections=(
            TopicSection(heading="1. 短标题", text="正文。", source_message_ids=(MESSAGE_ID,)),
            TopicSection(heading="2. English 2026 与中文标点：这是一个需要自然换行的较长 Idea 标题", text="正文。", source_message_ids=(MESSAGE_ID,)),
        )),
    })
    version = original.current_version.model_copy(update={"content": content})
    docx = Document(BytesIO(render_docx(version, original.source_evidence, source_mode="body_only")))

    assert [paragraph.style.name for paragraph in docx.paragraphs if paragraph.text.startswith(("1. 短", "2. English"))] == [
        "Heading 1", "Heading 1",
    ]
    assert docx.styles["Heading 1"].font.size > docx.styles["Normal"].font.size
    for style_name in ("Normal", "Title", "Heading 1", "Heading 2", "Title Char", "Heading 1 Char", "Heading 2 Char"):
        fonts = docx.styles[style_name]._element.rPr.rFonts
        assert {fonts.get(qn(f"w:{attribute}")) for attribute in ("ascii", "hAnsi", "eastAsia", "cs")} == {"Hiragino Sans GB"}
        assert not any(fonts.get(qn(f"w:{attribute}")) for attribute in
                       ("asciiTheme", "hAnsiTheme", "eastAsiaTheme", "cstheme"))


def test_legacy_user_edit_uses_saved_title_and_never_guesses_body_headings() -> None:
    version = _edited_version("已保存标题\n\n看起来像章节\n\n普通正文")

    markdown = render_markdown(version, (), source_mode="body_only").decode()
    assert markdown == "# 已保存标题\n\n看起来像章节\n\n普通正文\n"
    document = Document(BytesIO(render_docx(version, (), source_mode="body_only")))
    assert [(paragraph.style.name, paragraph.text) for paragraph in document.paragraphs] == [
        ("Title", "已保存标题"),
        ("Normal", "看起来像章节"),
        ("Normal", "普通正文"),
    ]


class DocumentStub:
    def get_document(self, document_id: str, *, session_id: str):
        assert document_id == "document:p3" and session_id == "session-a"
        return _document()


class ArtifactRepositoryStub:
    def __init__(self): self.rows = []
    def record_document_artifact(self, **row):
        self.rows.append(row)
        return {**row, "sha256": row.pop("sha256_digest")}


def test_export_writes_private_atomic_nonoverwriting_safe_filenames(tmp_path: Path) -> None:
    repository_root = tmp_path / "repo"; repository_root.mkdir()
    output = tmp_path / "approved"; output.mkdir(mode=0o700)
    root = ApprovedOutputRoot.approve(output, root_id="approved-root", repository_root=repository_root)
    repository = ArtifactRepositoryStub()
    service = DocumentExportService(document_service=DocumentStub(), repository=repository,
        approved_roots={"approved-root": root}, clock=lambda: NOW)

    first = service.export("document:p3", version_id="document-version:p3", format="markdown",
        source_mode="body_only", output_root_id="approved-root", session_id="session-a")
    second = service.export("document:p3", version_id="document-version:p3", format="docx",
        source_mode="with_sources", output_root_id="approved-root", session_id="session-a")

    assert first.filename != second.filename
    assert all("/" not in item.filename and ".." not in item.filename for item in (first, second))
    for item in (first, second):
        path = output / "documents" / item.filename
        assert path.is_file() and not path.is_symlink()
        assert stat.S_IMODE(path.stat().st_mode) == 0o600
        assert item.byte_size == path.stat().st_size
        assert str(output) not in item.model_dump_json()
    assert stat.S_IMODE((output / "documents").stat().st_mode) == 0o700
    assert len(repository.rows) == 2
